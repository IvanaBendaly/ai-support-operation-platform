import hashlib
import os
import subprocess
import sys
from datetime import UTC, datetime, timedelta
from io import BytesIO
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from alembic import command
from alembic.config import Config
from docx import Document as WordDocument
from sqlalchemy import event, select

from northstar.config import Settings
from northstar.db import make_engine, session_factory
from northstar.retrieval.extraction import IngestionError, extract
from northstar.retrieval.models import Document, Section
from northstar.retrieval.service import KnowledgeService, describe, retrieval_sections
from northstar.retrieval.worker import claim, convert_isolated, finish, run_once


def docx():
    document = WordDocument()
    document.add_heading("Employee leave", level=1)
    document.add_paragraph("Do NOT approve leave without manager approval. Allow 20 days.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Plan", "Days"
    table.cell(1, 0).text, table.cell(1, 1).text = "Standard", "20"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def enqueue(sessions, policy="leave", data=b"A readable policy", filename="a.txt", **kwargs):
    return KnowledgeService(sessions).ingest(policy, filename, data, defer=True, **kwargs)[
        "document"
    ]


def test_real_docx_headings_tables_and_worker_persistence(business):
    sessions, _ = business
    result = enqueue(sessions, data=docx(), filename="leave.docx")
    assert result["status"] == "QUEUED" and result["sections"] == []
    with sessions() as session:
        assert not retrieval_sections(session)
    assert run_once(sessions)
    with sessions() as session:
        document = session.get(Document, result["id"])
        assert document.status == "READY" and document.attempts == 1
        sections = retrieval_sections(session)
        text = "\n".join(s.text for s in sections)
        assert "Do NOT approve" in text and "Standard" in text and "20" in text
        assert any("heading:Employee leave" in s.source for s in sections)
        assert all("item:" in s.source and len(s.text) <= 2000 for s in sections)


def test_byte_duplicate_does_not_extract_or_enqueue_again(business):
    sessions, _ = business
    first = enqueue(sessions)
    with patch(
        "northstar.retrieval.service.extract", side_effect=AssertionError("must not extract")
    ):
        repeated = KnowledgeService(sessions).ingest("other", "renamed.txt", b"A readable policy")
    assert repeated["duplicate"] and repeated["document"]["id"] == first["id"]
    assert run_once(sessions)
    assert not run_once(sessions)


def test_normalized_duplicate_is_not_published(business):
    sessions, _ = business
    first = enqueue(sessions, data=b"Identical text\r\n")
    assert run_once(sessions)
    second = enqueue(sessions, policy="other", data=b"Identical text\n")
    assert run_once(sessions)
    with sessions() as session:
        duplicate = session.get(Document, second["id"])
        assert duplicate.status == "DUPLICATE" and duplicate.duplicate_of == first["id"]
        assert duplicate.content_hash is None
        assert len(retrieval_sections(session)) == 1
        assert not session.scalars(select(Section).where(Section.document_id == duplicate.id)).all()


def test_pending_and_failed_replacement_preserve_current_policy(business):
    sessions, _ = business
    first = enqueue(sessions)
    run_once(sessions)
    replacement = enqueue(sessions, filename="bad.docx", data=b"bad", replace=True)
    with sessions() as session:
        assert retrieval_sections(session)[0].document_id == first["id"]
    run_once(sessions)
    with sessions() as session:
        assert session.get(Document, replacement["id"]).error == "corrupt_docx"
        assert retrieval_sections(session)[0].document_id == first["id"]


def test_claim_expiry_fences_old_worker_and_bounds_attempts(business):
    sessions, _ = business
    document = enqueue(sessions)
    first = claim(sessions)
    assert claim(sessions) is None
    for _ in range(2):
        with sessions.begin() as session:
            session.get(Document, document["id"]).lease_expires_at = datetime.now(UTC) - timedelta(
                seconds=1
            )
        current = claim(sessions)
        assert current["token"] != first["token"]
    assert not finish(sessions, first, extract("a.txt", b"Stale text"))
    with sessions.begin() as session:
        session.get(Document, document["id"]).lease_expires_at = datetime.now(UTC) - timedelta(
            seconds=1
        )
    assert claim(sessions)["token"] is None
    with sessions() as session:
        failed = session.get(Document, document["id"])
        assert failed.status == "FAILED" and failed.error == "extraction_attempts_exhausted"
        assert failed.attempts == 3
    assert claim(sessions) is None


def test_partial_section_write_rolls_back_and_can_recover(business):
    sessions, _ = business
    document = enqueue(sessions)
    job = claim(sessions)
    result = extract("a.txt", b"Valid policy")

    def broken(*args):
        raise RuntimeError("injected write failure")

    event.listen(Section, "before_insert", broken)
    try:
        with pytest.raises(RuntimeError):
            finish(sessions, job, result)
    finally:
        event.remove(Section, "before_insert", broken)
    with sessions() as session:
        assert session.get(Document, document["id"]).status == "PROCESSING"
        assert not session.scalars(select(Section)).all()
    assert finish(sessions, job, result)


def test_conversion_timeout_is_bounded_and_inspectable(business):
    sessions, _ = business
    document = enqueue(sessions)
    with patch(
        "northstar.retrieval.worker.subprocess.run",
        side_effect=subprocess.TimeoutExpired("extract", 120),
    ):
        assert run_once(sessions)
    with sessions() as session:
        assert session.get(Document, document["id"]).error == "extraction_timeout"
        assert not retrieval_sections(session)
    # Actual process timeout, not only simulated exception classification.
    assert convert_isolated({"data": b"text", "filename": "a.txt"}, timeout=0.001) == {
        "error": "extraction_timeout"
    }


def test_docx_expansion_limit_and_corrupt_file():
    with pytest.raises(IngestionError, match="corrupt_docx"):
        extract("bad.docx", b"not a zip")
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * (26 * 1024 * 1024))
    with pytest.raises(IngestionError, match="docx_expanded_size_limit"):
        extract("large.docx", output.getvalue())


def test_worker_cli_reads_persisted_queue_after_restart(tmp_path, monkeypatch):
    url = f"sqlite+pysqlite:///{tmp_path / 'worker.db'}"
    monkeypatch.setenv("NORTHSTAR_DATABASE_URL", url)
    monkeypatch.setenv("NORTHSTAR_ENVIRONMENT", "test")
    command.upgrade(Config("alembic.ini"), "head")
    engine = make_engine(Settings(_env_file=None, database_url=url, environment="test"))
    try:
        sessions = session_factory(engine)
        document = enqueue(sessions, filename="--policy.txt")
        subprocess.run(
            [sys.executable, "-m", "northstar.retrieval.worker", "--once"],
            check=True,
            timeout=30,
            env=os.environ.copy(),
        )
        with sessions() as session:
            result = describe(session, session.get(Document, document["id"]), include_sections=True)
            assert result["status"] == "READY" and result["sections"]
            assert result["original_hash"] == hashlib.sha256(b"A readable policy").hexdigest()
        # Batch migration preserves existing originals and section relationships.
        command.downgrade(Config("alembic.ini"), "0006")
        command.upgrade(Config("alembic.ini"), "head")
        with sessions() as session:
            restored = session.get(Document, document["id"])
            assert restored.original == b"A readable policy" and restored.status == "READY"
            assert restored.attempts == 0
            assert retrieval_sections(session)[0].document_id == restored.id
    finally:
        engine.dispose()


@pytest.mark.postgres
def test_postgres_worker_claims_are_exclusive(postgres_business):
    from concurrent.futures import ThreadPoolExecutor

    sessions, _ = postgres_business
    enqueue(sessions)
    with ThreadPoolExecutor(max_workers=2) as pool:
        jobs = list(pool.map(lambda _: claim(sessions), range(2)))
    claimed = [j for j in jobs if j]
    assert len(claimed) == 1
    assert finish(sessions, claimed[0], extract("a.txt", b"A readable policy"))
